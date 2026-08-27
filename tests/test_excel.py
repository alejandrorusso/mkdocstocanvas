"""Tests for Excel rendering (processing/excel.py)."""

from pathlib import Path

from openpyxl import Workbook

from mkdocstocanvas.processing.excel import (
    apply_tint,
    color_to_hex,
    get_theme_colors,
    process_excel_macros,
    render_excel_sheet_to_html,
)


class TestApplyTint:
    def test_no_tint(self):
        assert apply_tint("FF0000", 0.0) == "FF0000"

    def test_positive_tint_lightens(self):
        assert apply_tint("000000", 1.0) == "FFFFFF"

    def test_negative_tint_darkens(self):
        assert apply_tint("FFFFFF", -1.0) == "000000"

    def test_clamped_to_byte_range(self):
        assert apply_tint("FFFFFF", 0.5) == "FFFFFF"


class TestGetThemeColors:
    def test_fallback_on_malformed_theme(self):
        workbook = type("FakeWorkbook", (), {"loaded_theme": b"<not-xml"})()
        colors = get_theme_colors(workbook)
        assert colors[0] == "FFFFFF"
        assert colors[4] == "4472C4"

    def test_workbook_without_theme(self):
        workbook = type("FakeWorkbook", (), {"loaded_theme": None})()
        colors = get_theme_colors(workbook)
        assert len(colors) == 10


class TestColorToHex:
    def test_none_color(self):
        assert color_to_hex(None, {}) is None

    def test_rgb_color(self, monkeypatch):
        # openpyxl returns 8-digit ARGB strings; the leading alpha is stripped
        color = type("C", (), {"rgb": "00FF0000"})()
        assert color_to_hex(color, {}) == "FF0000"

    def test_weird_openpyxl_string_is_ignored(self):
        # openpyxl sometimes returns a descriptive error string instead of a color
        color = type("C", (), {"rgb": "Values must be of type <class 'str'>"})()
        assert color_to_hex(color, {}) is None


class TestRenderExcelSheet:
    def test_missing_file_returns_error_html(self, tmp_path: Path):
        html = render_excel_sheet_to_html(tmp_path / "nope.xlsx")
        assert "Error" in html
        assert "nope.xlsx" in html

    def test_missing_sheet_returns_error_html(self, tmp_path: Path):
        wb = Workbook()
        wb.active["A1"] = "hi"
        file = tmp_path / "book.xlsx"
        wb.save(file)

        html = render_excel_sheet_to_html(file, sheet_name="No Such Sheet")
        assert "not found" in html

    def test_renders_table(self, tmp_path: Path):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Name"
        ws["A2"] = "Ada"
        file = tmp_path / "book.xlsx"
        wb.save(file)

        html = render_excel_sheet_to_html(file)
        assert "<table" in html
        assert "<th" in html and "Name" in html
        assert "<td" in html and "Ada" in html

    def test_blank_rows_skipped(self, tmp_path: Path):
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Header"
        ws["A3"] = "Value"  # row 2 left blank
        file = tmp_path / "book.xlsx"
        wb.save(file)

        html = render_excel_sheet_to_html(file)
        assert html.count("<tr>") == 2


class TestProcessExcelMacros:
    def test_macro_replaced_with_table(self, tmp_path: Path):
        wb = Workbook()
        wb.active["A1"] = "Cell"
        xlsx = tmp_path / "data.xlsx"
        wb.save(xlsx)

        md = tmp_path / "page.md"
        md.write_text(f"Before\n{{{{ render_excel_sheet('./{xlsx.name}') }}}}\nAfter")
        result = process_excel_macros(md.read_text(), md)

        assert "{{ render_excel_sheet" not in result
        assert "<table" in result
        assert "Before" in result and "After" in result

    def test_relative_path_resolved_from_markdown_dir(self, tmp_path: Path):
        wb = Workbook()
        wb.active["A1"] = "X"
        wb.save(tmp_path / "data.xlsx")

        sub = tmp_path / "sub"
        sub.mkdir()
        md = sub / "page.md"
        md.write_text("{{ render_excel_sheet('../data.xlsx') }}")

        result = process_excel_macros(md.read_text(), md)
        assert "<table" in result

    def test_text_without_macros_unchanged(self, tmp_path: Path):
        md = tmp_path / "page.md"
        md.write_text("No macros here")
        assert process_excel_macros("No macros here", md) == "No macros here"
