import re
from pathlib import Path
from typing import Dict, Optional

try:
    from openpyxl import load_workbook
    from openpyxl.styles.colors import COLOR_INDEX

    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


def apply_tint(rgb_hex: str, tint: float) -> str:
    """
    Apply a tint value to an RGB color.

    Tint values:
    - Positive tint: lightens the color (mixes with white)
    - Negative tint: darkens the color (mixes with black)
    - 0: no change
    """
    rgb_hex = rgb_hex.lstrip("#")

    r = int(rgb_hex[0:2], 16)
    g = int(rgb_hex[2:4], 16)
    b = int(rgb_hex[4:6], 16)

    if tint < 0:
        r = int(r * (1 + tint))
        g = int(g * (1 + tint))
        b = int(b * (1 + tint))
    else:
        r = int(r + (255 - r) * tint)
        g = int(g + (255 - g) * tint)
        b = int(b + (255 - b) * tint)

    r = max(0, min(255, r))
    g = max(0, min(255, g))
    b = max(0, min(255, b))

    return f"{r:02X}{g:02X}{b:02X}"


def get_theme_colors(workbook) -> Dict[int, str]:
    """
    Extract theme colors from workbook.
    Returns a dictionary mapping theme index to RGB hex color.
    """
    default_theme_colors = {
        0: "FFFFFF",  # White (background 1)
        1: "000000",  # Black (text 1)
        2: "E7E6E6",  # Light Gray (background 2)
        3: "44546A",  # Dark Blue Gray (text 2)
        4: "4472C4",  # Blue (accent 1)
        5: "ED7D31",  # Orange (accent 2)
        6: "A5A5A5",  # Gray (accent 3)
        7: "FFC000",  # Gold (accent 4)
        8: "5B9BD5",  # Light Blue (accent 5)
        9: "70AD47",  # Green (accent 6)
    }

    try:
        if hasattr(workbook, "loaded_theme") and workbook.loaded_theme:
            theme_xml = workbook.loaded_theme.decode("utf-8")

            import xml.etree.ElementTree as ET

            root = ET.fromstring(theme_xml)
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            color_scheme = root.find(".//a:clrScheme", ns)

            if color_scheme:
                theme_colors = {}
                color_map = {
                    "a:lt1": 0,
                    "a:dk1": 1,
                    "a:lt2": 2,
                    "a:dk2": 3,
                    "a:accent1": 4,
                    "a:accent2": 5,
                    "a:accent3": 6,
                    "a:accent4": 7,
                    "a:accent5": 8,
                    "a:accent6": 9,
                }

                for color_name, index in color_map.items():
                    color_elem = color_scheme.find(color_name, ns)
                    if color_elem:
                        srgb = color_elem.find(".//a:srgbClr", ns)
                        if srgb is not None and "val" in srgb.attrib:
                            theme_colors[index] = srgb.attrib["val"]
                            continue

                        sysclr = color_elem.find(".//a:sysClr", ns)
                        if sysclr is not None and "lastClr" in sysclr.attrib:
                            theme_colors[index] = sysclr.attrib["lastClr"]
                            continue

                if len(theme_colors) >= 6:
                    return theme_colors
    except Exception:
        pass

    return default_theme_colors


def color_to_hex(color_obj, theme_colors: Dict[int, str]) -> Optional[str]:
    """Convert an openpyxl Color object to hex RGB string."""
    if not color_obj:
        return None

    try:
        if hasattr(color_obj, "rgb") and color_obj.rgb:
            rgb_str = str(color_obj.rgb)
            if (
                rgb_str
                and "must be of type" not in rgb_str
                and rgb_str not in ["00000000", "FFFFFFFF", "0", "ffffffff"]
            ):
                if len(rgb_str) == 8:
                    return rgb_str[2:]
                elif len(rgb_str) == 6:
                    return rgb_str

        if hasattr(color_obj, "theme"):
            theme_val = color_obj.theme
            if isinstance(theme_val, int) and theme_val in theme_colors:
                base_color = theme_colors[theme_val]
                if hasattr(color_obj, "tint") and color_obj.tint:
                    return apply_tint(base_color, color_obj.tint)
                return base_color

        if hasattr(color_obj, "indexed"):
            idx_val = color_obj.indexed
            if isinstance(idx_val, int) and 0 <= idx_val < len(COLOR_INDEX):
                return COLOR_INDEX[idx_val].replace("#", "")

    except Exception:
        pass

    return None


def render_excel_sheet_to_html(excel_path: Path, sheet_name: str = None) -> str:
    """
    Render an Excel sheet as an HTML table, preserving cell colors.

    Args:
        excel_path: Path to the Excel file
        sheet_name: Name of the sheet to render (if None, uses first sheet)

    Returns:
        HTML table string
    """
    if not OPENPYXL_AVAILABLE:
        return "<p><em>Note: Excel sheet rendering not available (openpyxl not installed)</em></p>"

    if not excel_path.exists():
        return f"<p><em>Error: Excel file not found: {excel_path}</em></p>"

    try:
        workbook = load_workbook(excel_path, data_only=False)

        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                return f'<p><em>Error: Sheet "{sheet_name}" not found in {excel_path}</em></p>'
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.active

        theme_colors = get_theme_colors(workbook)

        html_parts = ['<div style="overflow-x: auto; margin: 1.5em 0;">']
        html_parts.append(
            '<table style="border-collapse: collapse; width: 100%; border: 1px solid #ddd;">'
        )

        first_row = True
        for row in sheet.iter_rows():
            if all(cell.value is None or str(cell.value).strip() == "" for cell in row):
                continue

            html_parts.append("<tr>")
            for cell in row:
                cell_value = str(cell.value) if cell.value is not None else ""

                bg_color = "white"
                text_color = "#000000"

                if cell.fill and cell.fill.fill_type and cell.fill.fill_type != "none":
                    if cell.fill.start_color:
                        bg_hex = color_to_hex(cell.fill.start_color, theme_colors)
                        if bg_hex:
                            bg_color = f"#{bg_hex}"

                if cell.font and cell.font.color:
                    text_hex = color_to_hex(cell.font.color, theme_colors)
                    if text_hex and text_hex.upper() not in [
                        "00000000",
                        "FFFFFFFF",
                        "FFFFFF",
                        "000000",
                    ]:
                        text_color = f"#{text_hex}"
                    if text_color == bg_color:
                        text_color = "#000000"

                is_bold = cell.font and cell.font.bold
                font_weight = "bold" if is_bold else "normal"

                if first_row:
                    header_bg = bg_color if bg_color != "white" else "#1976d2"
                    header_text = text_color if bg_color != "white" else "white"
                    header_weight = font_weight if font_weight == "bold" else "600"

                    html_parts.append(
                        f'<th style="background-color: {header_bg}; color: {header_text}; '
                        f"padding: 12px 16px; text-align: left; font-weight: {header_weight}; "
                        f'border: 1px solid #ddd;">{cell_value}</th>'
                    )
                else:
                    html_parts.append(
                        f'<td style="padding: 12px 16px; border: 1px solid #ddd; '
                        f"background-color: {bg_color}; color: {text_color}; "
                        f'font-weight: {font_weight};">{cell_value}</td>'
                    )
            html_parts.append("</tr>")
            first_row = False

        html_parts.append("</table>")
        html_parts.append("</div>")

        return "".join(html_parts)

    except Exception as e:
        return f"<p><em>Error rendering Excel sheet: {str(e)}</em></p>"


def process_excel_macros(content: str, markdown_file_path: Path) -> str:
    """
    Process Excel rendering macros in markdown content.
    Replaces {{ render_excel_sheet('./path/to/file.xlsx', 'SheetName') }} with rendered HTML tables.
    """
    markdown_dir = markdown_file_path.resolve().parent

    pattern = r'\{\{\s*render_excel_sheet\([\'"]([^\'"]+)[\'"](?:,\s*[\'"]([^\'"]+)[\'"])?\)\s*\}\}'

    def replace_excel_macro(match):
        excel_path_str = match.group(1)
        sheet_name = match.group(2) if match.group(2) else None

        excel_path = Path(excel_path_str)

        if not excel_path.is_absolute():
            excel_path = markdown_dir / excel_path

        excel_path = excel_path.resolve()

        return render_excel_sheet_to_html(excel_path, sheet_name)

    return re.sub(pattern, replace_excel_macro, content)
